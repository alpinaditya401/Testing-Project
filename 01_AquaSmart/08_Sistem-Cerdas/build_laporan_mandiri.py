"""Create the assignment report from executed notebook evidence."""
from pathlib import Path
import json
import re
import nbformat
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = HERE / 'tugas_mandiri'
ev = json.loads((OUT / 'evaluasi_mandiri.json').read_text(encoding='utf-8'))
eda = json.loads((OUT / 'eda.json').read_text(encoding='utf-8'))
nb1 = nbformat.read(ROOT / 'notebooks/01_eda.ipynb', as_version=4)
nb2 = nbformat.read(ROOT / 'notebooks/02_baseline.ipynb', as_version=4)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('AquaSmart AIoT  |  Tugas Mandiri Hari 1  |  ').font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def para(text):
    p = doc.add_paragraph()
    for i, part in enumerate(re.split(r'\*\*(.*?)\*\*', text.replace('`', ''))):
        p.add_run(part).bold = i % 2 == 1
    return p


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = True
    for c, value in zip(t.rows[0].cells, headers):
        c.text = str(value)
        for run in c.paragraphs[0].runs:
            run.bold = True
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'EDEFEF')
        c._tc.get_or_add_tcPr().append(shading)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c, value in zip(t.add_row().cells, row):
            c.text = str(value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(5)
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)


doc.add_paragraph('Tugas Mandiri Hari 1\nAquaSmart AIoT', 'Title')
doc.add_paragraph('Fondasi Data dan Machine Learning untuk Project Capstone', 'Subtitle')
para('Mata kuliah Python untuk Sistem Cerdas\nD3 Teknik Informatika Kampus Kabupaten Madiun\nSekolah Vokasi Universitas Sebelas Maret')
table(['Anggota tim', 'NIM'], [['Alpin Aditya Pratama', 'V3925004'], ['Dimas Aryo Sejati', 'V3925022']])
para('Laporan ini menjawab Soal 1–5 pada MODUL_PRAKTIKUM1 menggunakan dataset kolam yang telah tersedia di Testing-Project. Praktikum triase yang sudah dikerjakan Dimas tetap menjadi bagian terpisah. Project tugas mandiri adalah AquaSmart AIoT.')
para('Eksperimen memprediksi status proksi pH dan suhu untuk interval lima menit berikutnya. Pada holdout waktu, Random Forest memperoleh macro F1 0,8951, di bawah aturan ambang terakhir 0,9158. Hasil ini mendukung penggunaan baseline non-AI sebagai pembanding awal; belum ada bukti bahwa model ML perlu menggantikannya.')
para('**Batas hasil.** Label dibuat dari aturan aplikasi, bukan penilaian ahli. Dataset tidak memiliki kekeruhan sehingga hasil bukan status kualitas air lengkap. Evaluasi menggunakan satu sumber kolam dan belum merupakan validasi perangkat AquaSmart di lapangan.')
doc.add_heading('Berkas pelaksanaan', 2)
para('Kode dan output lengkap berada di notebooks/01_eda.ipynb serta notebooks/02_baseline.ipynb pada akar repository. Data bersih, model tersimpan, tabel hasil dan tangkapan layar lima soal berada di 01_AquaSmart/08_Sistem-Cerdas/tugas_mandiri/.')
para('**Kontribusi Git.** Modul mensyaratkan commit dari kedua anggota. Syarat ini belum diklaim selesai; setiap anggota harus meninjau dan melakukan commit sendiri. Tidak dibuat commit dengan identitas anggota lain.')

page('Soal 1 Project framing')
framing = nb1.cells[0].source.split('## Soal 1 Project framing\n')[1].split('### Sumber dan batas provenance')[0]
for block in framing.strip().split('\n\n')[:6]:
    para(block)
page('Soal 1 Batas lingkup dan sumber data')
for block in framing.strip().split('\n\n')[6:]:
    para(block)
doc.add_heading('Sumber dataset dan keterbatasan verifikasi', 2)
para('Boby Siswanto. A Simple Dataset of Aquaponic Fish Pond Water Quality Measurement using Internet of Things devices. Mendeley Data versi 2, 27 Maret 2023. DOI 10.17632/yd36bx6f8f.2. Lisensi pada halaman sumber: CC BY 4.0. https://data.mendeley.com/datasets/yd36bx6f8f/2')
para('README project mengatribusikan file lokal ke sumber tersebut. Halaman publik menyebut 118.286 baris terfilter, sedangkan file lokal berisi 505.730 baris. Kesetaraan file lokal dengan unduhan versi 2 belum diverifikasi. Eksperimen mengidentifikasi file lokal melalui SHA-256, bukan menyamakan kedua jumlah itu.')
para('SHA-256 file lokal: ' + eda['audit']['raw_sha256'])
para('Aturan dibaca langsung dari server/src/ThresholdRules.php: pH 6,5–8,5 dan suhu 25–30 °C. Ini adalah ambang default project, bukan klaim standar universal. TDS digunakan sebagai fitur tambahan; ambang kekeruhan tidak diterapkan pada TDS.')
para('Bukti output framing dan ambang yang dibaca dari kode: bukti/soal_1.png.')

page('Soal 2 EDA dataset project')
a = eda['audit']
table(['Pemeriksaan', 'Hasil eksekusi'], [
    ['Ukuran mentah', f"{a['raw_rows']:,} baris × {a['raw_columns']} kolom"],
    ['Kolom', 'id, created_date, water_pH, TDS, water_temp'],
    ['Duplikat semua kolom', a['exact_duplicates']],
    ['Tuple waktu dan nilai berulang yang digabung', a['repeated_measurements_collapsed']],
    ['Timestamp tidak valid', a['invalid_timestamp']],
    ['pH di luar 0 sampai 14', a['impossible_values']['ph']],
    ['Grid lima menit termasuk jeda', a['grid_bins']],
    ['Interval kosong', a['empty_bins']],
    ['Label proksi dalam ambang', eda['distribution']['jumlah']['Dalam ambang']],
    ['Label proksi di luar ambang', eda['distribution']['jumlah']['Di luar ambang']],
])
para('Tidak ada kategori substantif pada CSV. Kolom created_date dikonversi menjadi timestamp, sedangkan id adalah ID baris dan tidak digunakan sebagai fitur. df.info(), df.describe(), missing value per kolom serta lima baris pertama disertakan lengkap dalam notebook dan tangkapan layar soal_2.png.')
for insight in eda['insights']:
    para(insight)

page('Soal 2 Visualisasi hasil eksplorasi')
doc.add_picture(str(OUT / 'soal_2_grafik.png'), width=Inches(6.8))
para('Gambar 1. Empat grafik dari eksekusi notebook: distribusi label proksi, histogram pH, tren suhu harian, dan korelasi sensor. Grafik memakai interval yang tersedia, sedangkan bin kosong tetap dipertahankan dalam tabel untuk menjaga makna waktu.')
para('Alarm proksi dominan karena banyak pH atau suhu berada di luar rentang aplikasi. Hal ini bisa mencerminkan karakter kolam sumber atau drift sensor; dataset ini tidak cukup untuk menetapkan penyebab. Skor evaluasi tidak boleh diterjemahkan sebagai akurasi deteksi ikan sakit atau kekeruhan.')

page('Soal 3 Pembersihan dan fitur turunan')
for block in nb1.cells[8].source.split('\n', 1)[1].strip().split('\n\n'):
    para(block)
table(['Kelompok fitur', 'Makna dan ketersediaan'], [
    ['ph_lag1, tds_lag1, temp_lag1', 'Median sensor pada interval t−1 yang telah selesai'],
    ['ph_lag2, tds_lag2, temp_lag2', 'Median sensor pada interval t−2'],
    ['ph_mean12, tds_mean12, temp_mean12', 'Rata-rata maksimal 12 interval sebelum t, minimal 3 nilai'],
    ['ph_std12, tds_std12, temp_std12', 'Simpangan baku pada jendela historis yang sama'],
])
para('Data bersih disimpan sebagai data_bersih_5menit.csv. Fitur dan label proksi disimpan sebagai fitur_dan_target.csv. Missing value fitur sengaja masih terlihat pada file tersebut; imputasi dilakukan saat fit Pipeline. Kedua CSV tidak mengandung identitas akun atau kredensial.')

page('Soal 4 Baseline project')
para('Pembagian dilakukan menurut urutan waktu, sekitar 80 persen sampel awal untuk latih dan 20 persen akhir untuk uji. Sampel satu jam sebelum awal uji dikeluarkan dari latih. Semua baseline memakai sampel uji yang sama.')
table(['Bagian', 'Jumlah', 'Periode'], [[s['bagian'], s['jumlah'], s['awal'] + '\nsampai ' + s['akhir']] for s in ev['split']])
para('Baseline non-AI menerapkan ambang project ke pH dan suhu terakhir, kemudian mengasumsikan status tetap pada interval berikutnya. Dummy selalu memilih kelas mayoritas. Random Forest memakai 160 pohon, max_depth 10, min_samples_leaf 5, class_weight balanced, dan random_state 42.')
para('ColumnTransformer menjalankan SimpleImputer median dengan indikator missing serta StandardScaler pada 12 fitur numerik. Seluruh preprocessing dan Random Forest berada dalam Pipeline; fit hanya dijalankan pada data latih. Encoding tidak diperlukan karena tidak ada fitur kategorik.')
table(['Model', 'Macro F1', 'Balanced accuracy', 'Recall alarm', 'Accuracy'], [
    [r['model'], *[f"{r[k]:.4f}" for k in ['macro_f1', 'balanced_accuracy', 'recall_alarm', 'accuracy']]] for r in ev['results']])
para(ev['conclusion'])
para(f"TimeSeriesSplit lima fold pada data latih menghasilkan macro F1 rata-rata {ev['cv_macro_f1_mean']:.4f} dengan simpangan {ev['cv_macro_f1_std']:.4f}. Variasi antarperiode cukup besar; hasil satu holdout tidak cukup untuk menyatakan model stabil. Hasil per fold tersimpan pada cross_validation.csv.")
para('Fold validasi kelima hanya memuat kelas alarm. Macro F1 tetap dihitung pada dua label dengan zero_division 0, sehingga skor fold tersebut bukan pengukuran kemampuan membedakan dua kelas. Rata-rata lintas fold harus dibaca dengan batasan ini.')
para('Dummy mencapai recall alarm 1,0000 karena menandai semua sampel sebagai alarm, tetapi macro F1 hanya 0,4686. Ini menunjukkan mengapa recall dan accuracy tidak cukup bila dilaporkan sendiri.')

page('Soal 4 Pemeriksaan hasil prediksi')
doc.add_picture(str(OUT / 'soal_4_confusion_matrix.png'), width=Inches(6.8))
para('Gambar 2. Confusion matrix ketiga baseline pada 2.135 sampel uji. Dalam berarti pH dan suhu berada pada ambang aplikasi, bukan sertifikasi air aman; Luar berarti sedikitnya salah satu parameter melanggar ambang.')
para('Pipeline Random Forest disimpan pada pipeline_random_forest.joblib. Pengujian memuat ulang model dan membandingkan seluruh prediksi data uji menghasilkan prediksi identik. Gunakan file model yang dibuat sendiri dari notebook dan versi library yang tercatat.')
para('Bukti lengkap tabel split, skor validasi, perbandingan baseline, dan confusion matrix tersedia pada bukti/soal_4.png. Tidak dilakukan tuning setelah melihat hasil holdout. Integrasi model ke aplikasi atau aktuator tidak dilakukan pada tugas ini.')

page('Soal 5 Pemeriksaan data leakage')
for block in nb2.cells[7].source.split('\n', 1)[1].strip().split('\n\n'):
    para(block)
doc.add_heading('Verifikasi yang dijalankan', 2)
for text in [
    'Timestamp latih dan uji tidak beririsan dan batas latih berjarak lebih dari satu jam terhadap awal uji.',
    'Daftar fitur tepat 12 kolom historis; nilai sensor interval target dan target_proxy tidak masuk prediktor.',
    'Statistik median imputer identik dengan median data latih, bukan seluruh data.',
    'Mengubah sensor pada dan setelah waktu uji kausalitas tidak mengubah fitur pada atau sebelum waktu itu.',
    'Prediksi model setelah disimpan dan dimuat ulang identik dengan prediksi sebelum penyimpanan.'
]:
    para(text)

page('Lampiran bukti dan penyelesaian tim')
doc.add_picture(str(OUT / 'bukti/soal_5.png'), width=Inches(6.6))
para('Gambar 3. Tangkapan layar output audit fitur dan pemeriksaan otomatis dari notebook 02_baseline.ipynb.')
table(['Soal', 'Tangkapan layar output'], [[i, f'bukti/soal_{i}.png'] for i in range(1, 6)])
para('Dua notebook sudah dieksekusi dan menyimpan output. Tangkapan layar adalah hasil render HTML dari output notebook, bukan foto tampilan antarmuka Jupyter. File sumber HTML turut disertakan agar angka dapat ditelusuri.')
para('Yang masih harus dilakukan tim sebelum pengumpulan: tinjau hasil dan batasannya, lalu pastikan Alpin dan Dimas masing-masing melakukan commit menggunakan akun sendiri. Tidak ada klaim bahwa syarat dua commit sudah terpenuhi.')

doc.core_properties.title = 'Tugas Mandiri Hari 1 AquaSmart AIoT'
doc.core_properties.author = 'Alpin Aditya Pratama; Dimas Aryo Sejati'
doc.save(OUT / 'Tugas_Mandiri_Hari_1_AquaSmart.docx')
print(OUT / 'Tugas_Mandiri_Hari_1_AquaSmart.docx')
