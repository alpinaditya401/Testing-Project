"""Create the Hari 2 assignment report from the executed notebook 03 and its screenshots."""
from pathlib import Path
import re

import nbformat
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
OUT = HERE / 'tugas_mandiri_hari2'
BUKTI = OUT / 'bukti'
nb = nbformat.read(ROOT / 'notebooks/03_evaluasi_perbandingan.ipynb', as_version=4)
markdown_cells = [c.source for c in nb.cells if c.cell_type == 'markdown']

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    for border in list(style.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 2'].font.size = Pt(12)
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('AquaSmart AIoT  |  Tugas Mandiri Hari 2  |  ').font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def cell_starting(prefix):
    hits = [text for text in markdown_cells if text.lstrip().startswith(prefix)]
    assert len(hits) == 1, (prefix, len(hits))
    return hits[0]


def para(text, style=None):
    p = doc.add_paragraph(style=style)
    for part in re.split(r'(\*\*.*?\*\*|\*[^*\s][^*]*\*)', text.replace('`', '')):
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        elif part:
            p.add_run(part)
    return p


def markdown(text, skip=()):
    """Paragraphs, bullets, numbered items and ### titles of one notebook cell."""
    blocks, current = [], None
    for line in text.splitlines():
        s = line.strip()
        if not s or s == '---' or s in skip:
            current = None
        elif s.startswith('### '):
            blocks.append(['title', s[4:]])
            current = None
        elif s.startswith('- '):
            current = ['List Bullet', s[2:]]
            blocks.append(current)
        elif re.match(r'^\d+\. ', s):
            current = ['List Number', re.sub(r'^\d+\. ', '', s)]
            blocks.append(current)
        elif current is None:
            current = [None, s]
            blocks.append(current)
        else:
            current[1] += ' ' + s
    for style, content in blocks:
        if style == 'title':
            doc.add_heading(content, 2)
        else:
            para(content, style)


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = True
    for c, value in zip(t.rows[0].cells, headers):
        c.text = str(value)
        for run in c.paragraphs[0].runs:
            run.bold = True
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'EDEFEF')
        c._tc.get_or_add_tcPr().append(shading)
    for row in rows:
        for c, value in zip(t.add_row().cells, row):
            c.text = str(value)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.space_before = Pt(5)
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def page(title):
    # A break paragraph lands alone on a new page when the previous page is full.
    doc.add_heading(title, 1).paragraph_format.page_break_before = True


figures = 0


def figure(name, caption):
    """All output parts of one notebook cell, kept on one page with their caption."""
    global figures
    figures += 1
    parts = sorted(BUKTI.glob(f'{name}_p*_*.png'), key=lambda p: int(p.stem.split('_p')[-1].split('_')[0]))
    assert parts, name
    for path in parts:
        w, h = Image.open(path).size
        # Screenshots were taken at device scale 2, so w / 192 is the on-screen width in inches.
        width = min(6.8 if path.stem.endswith('_txt') else 6.0, w / 192)
        width = min(width, 8.2 * w / h)
        p = doc.add_paragraph()
        p.alignment = 1
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(width))
    para(f'Gambar {figures}. {caption}').runs[0].italic = True


setup_output = ''.join(o.get('text', '') for o in nb.cells[1].get('outputs', [])).strip().replace(' | ', ', ')
kesimpulan = next(text for text in markdown_cells if '### Kesimpulan Tugas Mandiri Hari 2' in text)
bullets = {line[2:].split(':', 1)[0]: line[2:] for line in re.sub(r'\n  +', ' ', kesimpulan).splitlines()
           if line.startswith('- ')}
batasan = bullets['Batasan'].split(':', 1)[1].strip()

doc.add_paragraph('Tugas Mandiri Hari 2\nAquaSmart AIoT', 'Title')
doc.add_paragraph('Evaluasi dan Perbandingan Model Machine Learning', 'Subtitle')
para('Mata kuliah Python untuk Sistem Cerdas\nD3 Teknik Informatika Kampus Kabupaten Madiun\nSekolah Vokasi Universitas Sebelas Maret')
table(['Anggota tim', 'NIM'], [['Alpin Aditya Pratama', 'V3925004'], ['Dimas Aryo Sejati', 'V3925022']])
para('Laporan ini menjawab Soal 1 sampai 3 pada MODUL_PRAKTIKUM2, Bagian 1 (Machine Learning dan Deep Learning), memakai data kolam hasil Tugas Mandiri Hari 1. Praktikum Hari 2 dengan data triase dari modul dikerjakan terpisah dan tidak dimuat di repository ini.')
para('**Hasil utama.** ' + bullets['Data uji'] + ' ' + bullets['Keputusan'])
para('**Batas hasil.** ' + batasan[0].upper() + batasan[1:])
doc.add_heading('Berkas pelaksanaan', 2)
para('Kode dan output lengkap berada di notebooks/03_evaluasi_perbandingan.ipynb pada akar repository. Tangkapan layar output dan laporan ini berada di 01_AquaSmart/08_Sistem-Cerdas/tugas_mandiri_hari2/. Model final disimpan di folder yang sama saat notebook dijalankan, tetapi tidak dimasukkan ke Git, sama seperti model Hari 1.')
para('Lingkungan eksekusi: ' + setup_output + '.')

page('Soal 1 Bagian, metrik, dan pembanding')
markdown(cell_starting('*Jawaban:*'), skip=('*Jawaban:*',))

page('Soal 2 Evaluasi dan perbandingan')
figure('s2_1', 'Output 2.1: data AquaSmart hasil Hari 1 dan pembagian latih/uji.')
figure('s2_2', 'Output 2.2: confusion matrix dan classification report baseline Hari 1 pada data uji.')
figure('s2_3', 'Output 2.3: perbandingan model dengan TimeSeriesSplit 5 fold pada data latih.')
figure('s2_4', 'Output 2.4: tuning Random Forest dengan GridSearchCV.')
figure('s2_5', 'Output 2.5: tabel perbandingan akhir (cross-validation) dan grafiknya.')
doc.add_heading('Model/konfigurasi yang dipilih dan alasannya', 2)
markdown(cell_starting('**Model/konfigurasi yang dipilih').split('\n\n', 1)[1])

page('Soal 3 Keputusan dan pemeriksaan')
figure('s3_1', 'Output 3.1: precision dan recall pada beberapa threshold (cross-validation, data latih).')
figure('s3_2', 'Output 3.2: pemeriksaan overfitting dengan Decision Tree berbagai kedalaman.')
figure('s3_3', 'Output 3.3: model final pada data uji yang dibuka satu kali, beserta confusion matrix.')
figure('s3_4', 'Output 3.4: ringkasan dan contoh kesalahan model final.')
figure('s3_5', 'Output 3.5: pemeriksaan dugaan penyebab kesalahan.')
doc.add_heading('Keputusan yang diambil dan contoh kesalahan', 2)
for prefix in ('**Alasan memilih threshold', '**Hasil pemeriksaan overfitting',
               '**Hasil uji (data uji dibuka', '**Contoh kesalahan dan dugaan'):
    markdown(cell_starting(prefix))
markdown(kesimpulan)

page('Lampiran bukti')
table(['Bukti', 'Berkas'], [['Seluruh output notebook 03 (HTML)', 'bukti/03_evaluasi_perbandingan.html'],
                            ['Soal 2, output 2.1 sampai 2.5', 'bukti/s2_*.png'],
                            ['Soal 3, output 3.1 sampai 3.5', 'bukti/s3_*.png']])
para('Tangkapan layar adalah hasil render HTML dari output notebook yang sudah dieksekusi, bukan foto tampilan antarmuka Jupyter. Berkas HTML disertakan agar setiap angka dapat ditelusuri ke output aslinya.')

doc.core_properties.title = 'Tugas Mandiri Hari 2 AquaSmart AIoT'
doc.core_properties.author = 'Alpin Aditya Pratama; Dimas Aryo Sejati'
doc.save(OUT / 'Tugas_Mandiri_Hari_2_AquaSmart.docx')
print(OUT / 'Tugas_Mandiri_Hari_2_AquaSmart.docx')
